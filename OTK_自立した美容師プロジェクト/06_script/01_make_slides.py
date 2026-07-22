# -*- coding: utf-8 -*-
"""
01_make_slides.py — slides_data.py から PowerPoint(16:9)を生成する。
出力: 01_slide/日本で一番自立した美容師が育つ会社.pptx
"""
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from slides_data import SLIDES, COLOR, FONT_SERIF, FONT_SANS, BRAND

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "01_slide", "日本で一番自立した美容師が育つ会社.pptx")

SW, SH = Inches(13.333), Inches(7.5)

def C(name):
    return RGBColor.from_string(COLOR[name])

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def set_bg(slide, color="bg"):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C(color)


def box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def para(tf, first=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    return p


def add_text(slide, x, y, w, h, text, *, font=FONT_SANS, size=16, color="white",
             bold=False, align=PP_ALIGN.CENTER, spacing=1.0, anchor=None,
             space_after=0):
    """複数行テキスト。text は str または (text, overrides) run タプルのリスト。"""
    tb, tf = box(slide, x, y, w, h)
    if anchor:
        tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else [text]
    for i, line in enumerate(lines):
        p = para(tf, first=(i == 0))
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        runs = [(line, {})] if isinstance(line, str) else line
        for rtext, ov in runs:
            r = p.add_run()
            r.text = rtext
            r.font.name = ov.get("font", font)
            r.font.size = Pt(ov.get("size", size))
            r.font.bold = ov.get("bold", bold)
            r.font.color.rgb = C(ov.get("color", color))
            _set_ea(r, ov.get("font", font))
    return tb


def _set_ea(run, font_name):
    """日本語(East Asian)フォントを明示指定して文字化けを防ぐ。"""
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea"):
        e = rPr.find(qn("a:" + tag))
        if e is None:
            e = rPr.makeelement(qn("a:" + tag), {})
            rPr.append(e)
        e.set("typeface", font_name)


def spaced(s):
    """ホテルサイン風レタースペーシング(文字間に空白挿入)。"""
    return " ".join(list(s.replace(" ", "  ")))


def card(slide, x, y, w, h, fill="card", line="line", radius=0.06):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = C(fill)
    sh.line.color.rgb = C(line)
    sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def chrome(slide, s):
    """全スライド共通の飾り: 上部キッカー、下部ページ番号。"""
    kicker = s.get("kicker", "")
    if kicker and s["layout"] not in ("title", "final"):
        add_text(slide, 0.5, 0.55, 12.333, 0.4, spaced(kicker),
                 font=FONT_SANS, size=11, color="gold", bold=False)
    add_text(slide, 11.6, 7.02, 1.35, 0.3, f"{s['no']:02d} / {len(SLIDES):02d}",
             font=FONT_SANS, size=9, color="gray", align=PP_ALIGN.RIGHT)
    if s["layout"] not in ("title", "final"):
        add_text(slide, 0.5, 7.02, 3.0, 0.3, "OTK", font=FONT_SERIF, size=9,
                 color="gold_dim", align=PP_ALIGN.LEFT)


def title_block(slide, s, y=1.15, size=30):
    add_text(slide, 0.9, y, 11.533, 1.7, s["title"], font=FONT_SERIF, size=size,
             color="white", spacing=1.2)


def note_block(slide, text, y=6.35):
    add_text(slide, 0.9, y, 11.533, 0.8, text, font=FONT_SANS, size=11,
             color="gray", spacing=1.25)


# ---------------- layouts ----------------

def l_title(slide, s):
    set_bg(slide, "bg")
    add_text(slide, 0.5, 2.0, 12.333, 0.45, spaced(s["kicker"]),
             font=FONT_SANS, size=13, color="gold")
    add_text(slide, 0.7, 2.75, 11.933, 2.2, s["title"], font=FONT_SERIF,
             size=40, color="white", spacing=1.25)
    add_text(slide, 0.5, 5.15, 12.333, 0.5, s["sub"], font=FONT_SANS,
             size=16, color="light")
    add_text(slide, 0.5, 5.95, 12.333, 0.4, "◆", font=FONT_SANS, size=9, color="gold_dim")


def l_question(slide, s):
    set_bg(slide, "bg")
    add_text(slide, 1.0, 2.35, 11.333, 2.9, s["title"], font=FONT_SERIF,
             size=33, color="white", spacing=1.35)


def l_statement(slide, s):
    set_bg(slide, "bg")
    nlines = s["title"].count("\n") + 1
    y = 1.7 if nlines >= 3 else 2.0
    add_text(slide, 0.9, y, 11.533, 0.6 + 0.75 * nlines, s["title"],
             font=FONT_SERIF, size=34, color="white", spacing=1.3)
    ly = y + 0.75 * nlines + 0.75
    for line in s.get("lines", []):
        add_text(slide, 1.4, ly, 10.533, 0.55, line, font=FONT_SANS, size=14.5,
                 color="light", spacing=1.3)
        ly += 0.62
    if s.get("gold_line"):
        add_text(slide, 0.9, ly + 0.25, 11.533, 0.7, s["gold_line"],
                 font=FONT_SERIF, size=21, color="gold")


def l_grid(slide, s):
    set_bg(slide, "bg")
    title_block(slide, s)
    items = s["items"]
    cols = 2 if len(items) == 4 else 3
    rows = (len(items) + cols - 1) // cols
    gw, gh, gap = (5.0, 1.15, 0.35) if cols == 2 else (3.7, 1.15, 0.3)
    total_w = cols * gw + (cols - 1) * gap
    x0 = (13.333 - total_w) / 2
    y0 = 3.15 if rows == 2 else 3.0
    if s["title"].count("\n") >= 1:
        y0 += 0.35
    for i, (head, desc) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (gw + gap), y0 + r * (gh + 0.3)
        card(slide, x, y, gw, gh)
        add_text(slide, x + 0.25, y + 0.17, gw - 0.5, 0.4, head, font=FONT_SANS,
                 size=15, bold=True, color="gold", align=PP_ALIGN.LEFT)
        add_text(slide, x + 0.25, y + 0.62, gw - 0.5, 0.4, desc, font=FONT_SANS,
                 size=11.5, color="light", align=PP_ALIGN.LEFT)
    if s.get("note"):
        note_block(slide, s["note"], y=y0 + rows * (gh + 0.3) + 0.25)


def l_stats(slide, s):
    set_bg(slide, "bg")
    title_block(slide, s, y=1.3)
    stats = s["stats"]
    gw, gap = 3.8, 0.35
    x0 = (13.333 - (gw * 3 + gap * 2)) / 2
    y0 = 3.1
    for i, (big, label) in enumerate(stats):
        x = x0 + i * (gw + gap)
        card(slide, x, y0, gw, 1.9)
        add_text(slide, x + 0.2, y0 + 0.45, gw - 0.4, 0.7, big, font=FONT_SERIF,
                 size=26, color="gold")
        add_text(slide, x + 0.25, y0 + 1.2, gw - 0.5, 0.55, label, font=FONT_SANS,
                 size=11.5, color="light", spacing=1.15)
    if s.get("note"):
        note_block(slide, s["note"], y=5.55)


def l_compare(slide, s):
    set_bg(slide, "bg")
    title_block(slide, s, y=1.25)
    for (head, big, tail), x in ((s["left"], 1.3), (s["right"], 7.23)):
        card(slide, x, 2.7, 4.8, 2.3)
        add_text(slide, x + 0.2, 2.98, 4.4, 0.4, head, font=FONT_SANS, size=13,
                 color="gray")
        add_text(slide, x + 0.2, 3.42, 4.4, 0.8, big, font=FONT_SERIF, size=34,
                 color="gold" if x > 5 else "white")
        add_text(slide, x + 0.2, 4.35, 4.4, 0.4, tail, font=FONT_SANS, size=12.5,
                 color="light")
    add_text(slide, 6.11, 3.55, 1.1, 0.6, "→", font=FONT_SANS, size=28, color="gold")
    note_block(slide, s["note"], y=5.55)


def l_twocol(slide, s):
    set_bg(slide, "bg")
    title_block(slide, s, y=1.2)
    coldef = [(s["left_head"], s["left_items"], "gray", "light", 1.3),
              (s["right_head"], s["right_items"], "gold", "white", 7.23)]
    for head, items, hcolor, icolor, x in coldef:
        card(slide, x, 2.35, 4.8, 3.35)
        add_text(slide, x + 0.35, 2.63, 4.1, 0.4, head, font=FONT_SANS, size=13.5,
                 bold=True, color=hcolor, align=PP_ALIGN.LEFT)
        yy = 3.18
        for it in items:
            add_text(slide, x + 0.35, yy, 4.1,
                     0.35, [("・ ", {"color": hcolor}), (it, {})],
                     font=FONT_SANS, size=12.5, color=icolor, align=PP_ALIGN.LEFT)
            yy += 0.47
    add_text(slide, 0.9, 6.05, 11.533, 0.5, s["note"], font=FONT_SERIF, size=15,
             color="gold", spacing=1.2)


def l_chips(slide, s):
    set_bg(slide, "bg")
    title_block(slide, s)
    chips = s["chips"]
    cols = 5 if len(chips) > 12 else (4 if len(chips) > 6 else len(chips))
    rows = (len(chips) + cols - 1) // cols
    gap = 0.25
    gw = min(2.5, (11.9 - (cols - 1) * gap) / cols)
    gh = 0.62
    y0 = 3.3 if s["title"].count("\n") else 2.9
    for r in range(rows):
        row_chips = chips[r * cols:(r + 1) * cols]
        row_w = len(row_chips) * gw + (len(row_chips) - 1) * gap
        x0 = (13.333 - row_w) / 2
        for c, chip in enumerate(row_chips):
            x, y = x0 + c * (gw + gap), y0 + r * (gh + 0.22)
            card(slide, x, y, gw, gh, radius=0.5)
            add_text(slide, x + 0.08, y + 0.14, gw - 0.16, 0.36, chip,
                     font=FONT_SANS, size=12.5, color="light")
    if s.get("note"):
        yn = y0 + rows * (gh + 0.22) + 0.25
        add_text(slide, 0.9, yn, 11.533, 0.8, s["note"], font=FONT_SANS, size=12,
                 color="gold", spacing=1.3)


def l_table(slide, s):
    set_bg(slide, "bg")
    add_text(slide, 0.9, 1.05, 11.533, 0.65, s["title"], font=FONT_SERIF, size=28,
             color="white")
    # 左: 3つの基準
    x, y = 1.0, 2.1
    for big, label in s["stats"]:
        card(slide, x, y, 5.6, 1.28)
        add_text(slide, x + 0.35, y + 0.2, 5.0, 0.55, big, font=FONT_SERIF,
                 size=22, color="gold", align=PP_ALIGN.LEFT)
        add_text(slide, x + 0.35, y + 0.78, 5.0, 0.4, label, font=FONT_SANS,
                 size=11.5, color="light", align=PP_ALIGN.LEFT)
        y += 1.48
    # 右: 月額↔年間 テーブル
    tx, ty, tw = 7.3, 2.1, 5.0
    add_text(slide, tx, ty, tw / 2, 0.4, s["table_head"][0], font=FONT_SANS,
             size=12, bold=True, color="gray")
    add_text(slide, tx + tw / 2, ty, tw / 2, 0.4, s["table_head"][1],
             font=FONT_SANS, size=12, bold=True, color="gray")
    ry = ty + 0.5
    for m, a, hl in s["table_rows"]:
        if hl:
            card(slide, tx - 0.1, ry - 0.06, tw + 0.2, 0.62, fill="card",
                 line="gold")
        color = "gold" if hl else "light"
        add_text(slide, tx, ry, tw / 2, 0.45, m, font=FONT_SANS,
                 size=15 if hl else 13.5, bold=hl, color=color)
        add_text(slide, tx + tw / 2, ry, tw / 2, 0.45, a, font=FONT_SANS,
                 size=15 if hl else 13.5, bold=hl, color=color)
        ry += 0.73
    note_block(slide, s["note"], y=6.6)


def l_savings(slide, s):
    set_bg(slide, "bg")
    title_block(slide, s, y=1.05)
    big, tail = s["big"]
    add_text(slide, 0.9, 2.0, 11.533, 0.85,
             [(big, {"font": FONT_SERIF, "size": 30, "color": "gold"}),
              ("  " + tail, {"size": 15, "color": "light"})])
    # 積立マイルストーン
    ms = s["milestones"]
    gw, gap = 2.15, 0.25
    x0 = (13.333 - (gw * len(ms) + gap * (len(ms) - 1))) / 2
    y0 = 3.15
    for i, (yrs, amt) in enumerate(ms):
        x = x0 + i * (gw + gap)
        card(slide, x, y0, gw, 1.25)
        add_text(slide, x, y0 + 0.18, gw, 0.4, yrs, font=FONT_SANS, size=12.5,
                 color="gray")
        add_text(slide, x, y0 + 0.6, gw, 0.5, amt, font=FONT_SERIF, size=17,
                 color="white")
    # 使いみち
    add_text(slide, 0.9, 4.75, 11.533, 0.4, "このお金がつくるのは、老後資金だけではなく「選択肢」",
             font=FONT_SANS, size=13, color="gold")
    add_text(slide, 0.9, 5.25, 11.533, 0.45, " ・ ".join(s["purposes"]),
             font=FONT_SANS, size=12.5, color="light")
    note_block(slide, s["note"], y=6.15)


def l_define(slide, s):
    set_bg(slide, "bg")
    add_text(slide, 0.9, 1.55, 11.533, 1.6, s["title"], font=FONT_SERIF, size=32,
             color="white", spacing=1.3)
    card(slide, 2.1, 3.95, 9.133, 1.75, fill="card", line="gold")
    add_text(slide, 2.5, 4.35, 8.333, 1.1, s["definition"], font=FONT_SERIF,
             size=19, color="gold", spacing=1.4)


def l_threelines(slide, s):
    set_bg(slide, "bg")
    title_block(slide, s, y=1.15)
    y = 3.1
    for i, pt_ in enumerate(s["points"], 1):
        add_text(slide, 2.2, y, 1.0, 0.6, f"{i:02d}", font=FONT_SERIF, size=22,
                 color="gold", align=PP_ALIGN.LEFT)
        add_text(slide, 3.3, y + 0.05, 8.5, 0.6, pt_, font=FONT_SANS, size=19,
                 color="white", align=PP_ALIGN.LEFT)
        y += 0.85
    note_block(slide, s["note"], y=6.1)


def l_cycle(slide, s):
    set_bg(slide, "bg")
    add_text(slide, 0.9, 1.0, 11.533, 0.6, s["title"], font=FONT_SERIF, size=28,
             color="white")
    nodes = list(zip(s["cycle"], s["cycle_sub"]))
    gw, gh, gap = 3.35, 1.15, 0.85
    x0 = (13.333 - (gw * 3 + gap * 2)) / 2
    ty, by = 2.15, 4.6
    pos = [(x0, ty), (x0 + gw + gap, ty), (x0 + 2 * (gw + gap), ty),
           (x0 + 2 * (gw + gap), by), (x0 + gw + gap, by), (x0, by)]
    for (name, sub), (x, y) in zip(nodes, pos):
        card(slide, x, y, gw, gh)
        add_text(slide, x + 0.15, y + 0.18, gw - 0.3, 0.42, name, font=FONT_SANS,
                 size=14.5, bold=True, color="gold")
        add_text(slide, x + 0.15, y + 0.63, gw - 0.3, 0.38, sub, font=FONT_SANS,
                 size=10.5, color="light")
    # 矢印(上段 →、右辺 ↓、下段 ←、左辺 ↑ で循環)
    ay = ty + 0.35
    for x in (x0 + gw, x0 + 2 * gw + gap):
        add_text(slide, x, ay, gap, 0.5, "→", font=FONT_SANS, size=20, color="gold")
    add_text(slide, x0 + 2 * (gw + gap) + gw / 2 - 0.25, ty + gh + 0.25, 0.5, 0.5,
             "↓", font=FONT_SANS, size=20, color="gold")
    ay = by + 0.35
    for x in (x0 + gw, x0 + 2 * gw + gap):
        add_text(slide, x, ay, gap, 0.5, "←", font=FONT_SANS, size=20, color="gold")
    add_text(slide, x0 + gw / 2 - 0.25, ty + gh + 0.25, 0.5, 0.5, "↑",
             font=FONT_SANS, size=20, color="gold")
    add_text(slide, 0.9, 6.15, 11.533, 0.5, "全員で稼いで、全員に還す。",
             font=FONT_SERIF, size=16, color="gold")


def l_final(slide, s):
    set_bg(slide, "bg")
    add_text(slide, 0.5, 0.95, 12.333, 0.4, spaced(s["kicker"]), font=FONT_SANS,
             size=11, color="gold")
    add_text(slide, 1.2, 1.85, 10.933, 1.9, s["message"], font=FONT_SERIF, size=20,
             color="light", spacing=1.6)
    add_text(slide, 0.9, 4.35, 11.533, 1.5, s["title"], font=FONT_SERIF, size=36,
             color="white", spacing=1.3)


LAYOUTS = dict(title=l_title, question=l_question, statement=l_statement,
               grid=l_grid, stats=l_stats, compare=l_compare, twocol=l_twocol,
               chips=l_chips, table=l_table, savings=l_savings, define=l_define,
               threelines=l_threelines, cycle=l_cycle, final=l_final)

for s in SLIDES:
    slide = prs.slides.add_slide(BLANK)
    set_bg(slide)
    LAYOUTS[s["layout"]](slide, s)
    chrome(slide, s)
    slide.notes_slide.notes_text_frame.text = "".join(s["narration"])

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("saved:", OUT)
